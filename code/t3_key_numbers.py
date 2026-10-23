#!/usr/bin/env python3
"""关键数字逐项核对：正文断言 → 图/表来源。"""
import sys, re, glob, os
sys.path.insert(0, "code")
import docx, fitz
from openpyxl import load_workbook

DOCX = "选题3_网络毒理学_AI骨质疏松/manuscript/04. Manuscript_选题3_Environmental_Pollution版.docx"
paras = [p.text.strip() for p in docx.Document(DOCX).paragraphs if p.text.strip()]
body = " ".join(paras[paras.index("1. Introduction"):paras.index("Funding")])

figs = {}
for f in sorted(glob.glob("投稿材料/选题3_正图/Figure*.pdf")):
    dd = fitz.open(f); figs[os.path.basename(f).replace("Figure", "Fig").replace(".pdf", "")] = dd[0].get_text().replace("\n", " "); dd.close()
sis = {}
for f in sorted(glob.glob("选题3_网络毒理学_AI骨质疏松/figures/S*.pdf")):
    dd = fitz.open(f); sis[os.path.basename(f).replace(".pdf", "")] = dd[0].get_text().replace("\n", " "); dd.close()

wb = load_workbook("投稿材料/选题3_附表.xlsx", data_only=True)
def sheet(n):
    for ws in wb:
        if ws.title.startswith(f"TableS{n}_"):
            vals = []
            for row in ws.iter_rows(values_only=True):
                for c in row:
                    if isinstance(c, (int, float)): vals.append(float(c))
                    elif isinstance(c, str):
                        for m in re.findall(r"-?\d+(?:\.\d+)?", c):
                            try: vals.append(float(m))
                            except: pass
            return vals
    return []

def has(vals, v):
    return any(abs(x - v) <= max(abs(v)*0.006, 1e-9) for x in vals)

def txt_has(text, v):
    for m in re.findall(r"-?\d+(?:\.\d+)?", text):
        try: x = float(m)
        except: continue
        if abs(x - v) <= max(abs(v)*0.006, 1e-9): return True
    return False

checks = [
    # (label, value, [figure keys], [table numbers], text-substring expected)
    ("ROR letrozole", 2.77, ["Fig1"], [2], "2.77"),
    ("ROR anastrozole", 4.08, ["Fig1"], [2], "4.08"),
    ("ROR exemestane", 3.61, ["Fig1"], [2], "3.61"),
    ("ROR letrozole fracture", 4.93, ["Fig1"], [2], "4.93"),
    ("ROR anastrozole fracture", 2.86, ["Fig1"], [2], "2.86"),
    ("tamoxifen ROR", 2.94, ["Fig1"], [2], "2.94"),
    ("E-value max", 9.33, [], [2, 15], "9.33"),
    ("E-value min", 4.98, [], [2, 15], "4.98"),
    ("letrozole co-reports", 290, [], [2, 15], "290"),
    ("anastrozole co-reports", 204, [], [2, 15], "204"),
    ("age n", 253, ["Fig1"], [15], "253"),
    ("sex n", 394, ["Fig1"], [15], "394"),
    ("female n", 390, ["Fig1"], [15], "390"),
    ("drug targets", 42, ["Fig2"], [1], "42"),
    ("disease genes", 3713, ["Fig2"], [1], "3,713"),
    ("core genes", 9, ["Fig2"], [1], "nine"),
    ("GO BP terms", 256, [], [3], "256"),
    ("GO CC terms", 15, [], [3], "15 CC"),
    ("GO MF terms", 33, [], [3], "33 MF"),
    ("DEG GSE35958", 1447, ["Fig3"], [10], "1,447"),
    ("DEG GSE13850", 28, ["Fig3"], [10], "28 DEGs"),
    ("DEG GSE56815", 391, ["Fig3"], [11], "391"),
    ("WGCNA modules", 13, ["Fig4"], [7], "13 modules"),
    ("M4 r", 0.363, ["Fig4"], [7], "0.363"),
    ("M5 r", 0.380, ["Fig4"], [7], "0.380"),
    ("M4 padj", 0.0066, ["Fig4"], [7], "0.0066"),
    ("M4 size", 1743, ["Fig4"], [7], "1,743"),
    ("M5 size", 931, ["Fig4"], [7], "931"),
    ("lambda min", 0.060, ["Fig5"], [9], "0.060"),
    ("lambda 1se", 0.146, ["Fig5"], [9], "0.146"),
    ("bootstrap %", 87.5, ["Fig5"], [15], "87.5"),
    ("AUC", 0.79, ["Fig5"], [15], "0.79"),
    ("FPKM VPS35", 12.97, ["Fig5"], [6], "12.97"),
    ("identity", 51.3, ["Fig7"], [14], "51.3"),
    ("RMSD all", 0.60, [], [12], "0.60"),
    ("RMSD pocket", 0.29, [], [12], "0.29"),
    ("redock RMSD", 0.32, [], [12], "0.32"),
    ("dock human let", 7.22, ["Fig6", "Fig7"], [14], "7.22"),
    ("dock human ana", 7.12, ["Fig6", "Fig7"], [14], "7.12"),
    ("dock fish let", 7.63, ["Fig7"], [14], "7.63"),
    ("dock fish ana", 7.20, ["Fig7"], [14], "7.20"),
    ("dock VPS35 let", 6.03, ["Fig6"], [14], "6.03"),
    ("dock VPS35 ana", 6.00, ["Fig6"], [14], "6.00"),
    ("BSA aromatase", 391, [], [12], "391"),
    ("BSA VPS35", 268, [], [12], "268"),
    ("vKO genes", 25, ["Fig6"], [10], "25 genes"),
    ("EAR exemestane 100", 0.704, ["Fig7"], [17], "0.70"),
    ("EAR letrozole 100", 0.176, ["Fig7"], [17], "0.18"),
    ("EAR anastrozole 100", 0.092, ["Fig7"], [17], "0.09"),
]

print(f"{'项目':<24}{'值':>8}  图  表  文  结论")
bad = []
for label, v, fkeys, tabs, needle in checks:
    f_ok = any(txt_has(figs[k], v) for k in fkeys if k in figs) if fkeys else None
    t_ok = any(has(sheet(n), v) for n in tabs) if tabs else None
    x_ok = needle in body
    ok = x_ok and (f_ok is not False) and (t_ok is not False)
    if not ok: bad.append((label, v, fkeys, tabs, f_ok, t_ok, x_ok))
    print(f"{label:<24}{v:>8}  {'Y' if f_ok else ('-' if f_ok is None else 'N')}  "
          f"{'Y' if t_ok else ('-' if t_ok is None else 'N')}  {'Y' if x_ok else 'N'}  {'OK' if ok else 'CHECK'}")

print("\n需人工确认：")
for b in bad: print("  ", b)
