"""T3 后处理：引文 () → []；生成 EES 双倍行距+行号版"""
import re, sys
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "选题3_网络毒理学_AI骨质疏松/manuscript/02. Manuscript_选题3_网络毒理学_AI骨质疏松_v2.docx"
EP = "选题3_网络毒理学_AI骨质疏松/manuscript/04. Manuscript_选题3_Environmental_Pollution版.docx"
EES = "选题3_网络毒理学_AI骨质疏松/manuscript/03. Manuscript_选题3_EES_双倍行距版.docx"

def brackets(doc):
    pat = re.compile(r"\((\d{1,3}(?:,\s*\d{1,3})*)\)")
    changed = 0
    for p in doc.paragraphs:
        for run in p.runs:
            if not run.text: continue
            new = pat.sub(lambda m: "[" + m.group(1) + "]", run.text)
            if new != run.text:
                run.text = new; changed += 1
    return changed

def typography(doc):
    """统一负号、乘号与科学计数法上标（可重复运行）"""
    import re
    n = 0
    for para in doc.paragraphs:
        for run in para.runs:
            t = run.text
            if not t:
                continue
            t2 = re.sub(r'(?<![\dA-Za-z])-(?=\d)', '\u2212', t)
            t2 = t2.replace("2x2", "2\u00d72")
            t2 = re.sub(r'10-([654])(?!\d)', lambda m: '10' + {'6': '\u207b\u2076', '5': '\u207b\u2075', '4': '\u207b\u2074'}[m.group(1)], t2)
            if t2 != t:
                run.text = t2
                n += 1
    return n


def main():
    d = docx.Document(SRC)
    n = brackets(d)
    d.save(SRC)
    print(f"引文格式 () → []：{n} 处")
    # 双倍行距 + 行号版
    d2 = docx.Document(SRC)
    for p in d2.paragraphs:
        if p.style.name == "Normal":
            p.paragraph_format.line_spacing = 2.0
    for sec in d2.sections:
        ln = OxmlElement('w:lnNumType')
        ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous')
        sec._sectPr.append(ln)
    typography(d2)
    d2.save(EES)
    print("EES 版（双倍行距+行号）重建")

    # EP 送审版（双倍行距 + 行号 + 参考文献双倍），并同步到 submission
    import shutil, os
    d3 = docx.Document(SRC)
    for p in d3.paragraphs:
        if (p.style.name == "Normal" or p.style.name.startswith("Bibliography")) and p.text.strip():
            p.paragraph_format.line_spacing = 2.0
    for sec in d3.sections:
        ln = OxmlElement('w:lnNumType')
        ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous')
        sec._sectPr.append(ln)
    typography(d3)
    d3.save(EP)
    sub = os.path.join(os.path.dirname(EP), "submission", "01_Manuscript_Environmental_Pollution.docx")
    if os.path.isdir(os.path.dirname(sub)):
        shutil.copy(EP, sub)
    print("EP 版（A4 送审）重建" + ("并同步 submission" if os.path.isdir(os.path.dirname(sub)) else ""))

if __name__ == "__main__":
    main()
